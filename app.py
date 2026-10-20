import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import os
import io
import hmac
import base64
import openpyxl
import requests

st.set_page_config(
    page_title="Sales & Distribution Dashboard",
    page_icon="logo.png",
    layout="wide"
)

st.markdown("""
    <style>
    div[data-testid="stMetric"] {
        background-color: color-mix(in srgb, currentColor 5%, transparent);
        border: 1px solid color-mix(in srgb, currentColor 15%, transparent);
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #007bff;
    }
    </style>
""", unsafe_allow_html=True)

SHARED_FILE = "shared_daily_sales.xlsx"

# قائمة المستخدمين والصلاحيات (افتراضية)
# الأفضل تعريفهم في .streamlit/secrets.toml تحت [users.<name>] بدل كتابتهم في الكود
DEFAULT_USERS = {
    "admin": {"password": "admin", "role": "admin", "line": "All"},
    "chc1": {"password": "123", "role": "user", "line": "Mash CHC 1"},
    "chc2": {"password": "123", "role": "user", "line": "Mash CHC 2"},
    "cns1": {"password": "123", "role": "user", "line": "Mash CNS 1"},
    "cns2": {"password": "123", "role": "user", "line": "Mash CNS 2"},
    "cns3": {"password": "123", "role": "user", "line": "Mash CNS 3"},
    "cns4": {"password": "123", "role": "user", "line": "Mash CNS 4"},
    "git1": {"password": "123", "role": "user", "line": "Mash GIT 1"},
    "git2": {"password": "123", "role": "user", "line": "Mash GIT 2"},
    "git3": {"password": "123", "role": "user", "line": "Mash GIT 3"},
    "cvm1": {"password": "123", "role": "user", "line": "Mash CVM 1"},
    "cvm2": {"password": "123", "role": "user", "line": "Mash CVM 2 New"},
    "ortho": {"password": "123", "role": "user", "line": "Mash Ortho"},
    "gyna": {"password": "123", "role": "user", "line": "Mash Gyna"},
    "dnu": {"password": "123", "role": "user", "line": "Mash DNU"},
    "oxa": {"password": "123", "role": "user", "line": "Mash Oxaleptal-Atten"},
    "derma": {"password": "123", "role": "user", "line": "Mash DERMA"},
    "primary1": {"password": "123", "role": "user", "line": "Mash Primary Care-1"},
    "primary2": {"password": "123", "role": "user", "line": "Mash Primary Care -2"},
    "medical": {"password": "123", "role": "user", "line": "Mash Medical Line -N"}
}


def load_users():
    """يقرأ المستخدمين من st.secrets لو موجودين، وإلا يستخدم DEFAULT_USERS."""
    try:
        if "users" in st.secrets:
            return {name: dict(cfg) for name, cfg in st.secrets["users"].items()}
    except Exception:
        pass
    return DEFAULT_USERS

USERS = load_users()

def check_credentials(username, password):
    user = USERS.get(username)
    if user is None:
        return None
    # مقارنة آمنة ضد timing attacks
    if hmac.compare_digest(str(user["password"]).encode("utf-8"), password.encode("utf-8")):
        return user
    return None

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.assigned_line = ""

# اسم العمود -> رقم العمود في الشيت
NUMERIC_COLS = {
    'Private_Sales_Qty': 3, 'Private_Sales_Val': 4,
    'Private_Target_Qty': 5, 'Private_Target_Val': 6,
    'Ibn Sina': 9, 'Pharma': 11, 'Ramco': 13, 'UCP': 15, 'MEC': 17,
    'Sofico': 19, 'ETC': 21, 'Direct': 23, 'Stores': 25,
    'Stock_Days': 42, 'Total_Stock_Qty': 45, 'Stock_Val_ex_Badr': 44,
    'Tender_Sales_Qty': 58, 'Tender_Sales_Val': 59,
    'Total_Sales_Qty': 60, 'Total_Sales_Val': 61,
    'Total_Target_Qty': 62, 'Total_Target_Val': 63,
}
DIST_NAMES = ['Ibn Sina', 'Pharma', 'Ramco', 'UCP', 'MEC', 'Sofico', 'ETC', 'Direct', 'Stores']

def format_currency(val):
    if abs(val) >= 1_000_000:
        return f"{val / 1_000_000:.2f}M EGP"
    return f"{val:,.0f} EGP"

# ---------- تخزين دائم للشيت على GitHub ----------
# Streamlit Cloud بيمسح أي ملف مرفوع عند كل restart، فبنحفظ نسخة في repo خاص.
# الإعداد من App settings > Secrets:
#   [github]
#   token  = "github_pat_..."     # Fine-grained token بصلاحية Contents: Read and write
#   repo   = "owner/repo-name"    # الأفضل repo خاص منفصل للداتا (من غير redeploy)
#   path   = "shared_daily_sales.xlsx"   # اختياري
#   branch = "main"                      # اختياري
def github_cfg():
    try:
        cfg = st.secrets["github"]
        if cfg.get("token") and cfg.get("repo"):
            return {
                "token": cfg["token"],
                "repo": cfg["repo"],
                "path": cfg.get("path", SHARED_FILE),
                "branch": cfg.get("branch", "main"),
            }
    except Exception:
        pass
    return None

def _gh_url(cfg):
    return f"https://api.github.com/repos/{cfg['repo']}/contents/{cfg['path']}"

def _gh_headers(cfg, raw=False):
    return {
        "Authorization": f"Bearer {cfg['token']}",
        "Accept": "application/vnd.github.raw+json" if raw else "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }

def push_to_github(content, uploader=""):
    """يرفع الشيت لـ GitHub. بيرجّع (نجح؟، رسالة)."""
    cfg = github_cfg()
    if cfg is None:
        return False, "GitHub مش مفعّل في الـ Secrets"
    try:
        url = _gh_url(cfg)
        r = requests.get(url, headers=_gh_headers(cfg), params={"ref": cfg["branch"]}, timeout=30)
        if r.status_code not in (200, 404):
            return False, f"GitHub رجّع خطأ {r.status_code}"
        payload = {
            "message": f"Update daily sales sheet (by {uploader or 'admin'})",
            "content": base64.b64encode(content).decode("ascii"),
            "branch": cfg["branch"],
        }
        if r.status_code == 200:
            payload["sha"] = r.json().get("sha")
        r = requests.put(url, headers=_gh_headers(cfg), json=payload, timeout=120)
        if r.status_code in (200, 201):
            return True, ""
        return False, f"GitHub {r.status_code}: {r.json().get('message', '')}"
    except Exception as e:
        return False, str(e)

def delete_from_github(deleter=""):
    """يمسح الشيت من GitHub. بيرجّع (نجح؟، رسالة). لو الملف مش موجود أصلاً يعتبرها نجاح."""
    cfg = github_cfg()
    if cfg is None:
        return True, ""
    try:
        url = _gh_url(cfg)
        r = requests.get(url, headers=_gh_headers(cfg), params={"ref": cfg["branch"]}, timeout=30)
        if r.status_code == 404:
            return True, ""
        if r.status_code != 200:
            return False, f"GitHub error {r.status_code}"
        payload = {
            "message": f"Delete daily sales sheet (by {deleter or 'admin'})",
            "sha": r.json().get("sha"),
            "branch": cfg["branch"],
        }
        r = requests.delete(url, headers=_gh_headers(cfg), json=payload, timeout=60)
        if r.status_code == 200:
            return True, ""
        return False, f"GitHub {r.status_code}: {r.json().get('message', '')}"
    except Exception as e:
        return False, str(e)

@st.cache_resource(show_spinner="Restoring latest sheet...", ttl=300)
def restore_shared_file():
    """لو الشيت مش موجود محلياً (بعد restart) بنسحبه من GitHub."""
    if os.path.exists(SHARED_FILE):
        return "local"
    cfg = github_cfg()
    if cfg is None:
        return "no-github"
    try:
        r = requests.get(_gh_url(cfg), headers=_gh_headers(cfg, raw=True),
                         params={"ref": cfg["branch"]}, timeout=120)
        if r.status_code == 200 and r.content:
            with open(SHARED_FILE, "wb") as f:
                f.write(r.content)
            return "restored"
        return f"missing ({r.status_code})"
    except Exception as e:
        return f"failed: {e}"

MONEY_COLS = ['Total_Sales_Val', 'Total_Target_Val', 'Private_Sales_Val', 'Tender_Sales_Val', 'Stock_Val_ex_Badr']

def style_numbers(frame):
    """يعرض الأرقام بفواصل آلاف (1,234,567) مع الحفاظ على الترتيب بالأرقام الحقيقية."""
    fmt = {}
    for c in frame.columns:
        if c in MONEY_COLS or c.endswith('_Qty'):
            fmt[c] = "{:,.0f}"
        elif c == 'Stock_Days':
            fmt[c] = "{:,.1f}"
    return frame.style.format(fmt, na_rep="-")

def resolve_data_file():
    return SHARED_FILE if os.path.exists(SHARED_FILE) else None

# الـ mtime بيتبعت كـ argument عشان الكاش يتجدد أوتوماتيك لما الأدمن يرفع شيت جديد
@st.cache_data(show_spinner=False)
def load_raw_sheet(file_path, mtime):
    return pd.read_excel(file_path, header=None)

def get_sheet_date(file_path):
    try:
        raw = load_raw_sheet(file_path, os.path.getmtime(file_path))
        value = raw.iloc[1, 2]
        return value if pd.notna(value) else None
    except Exception:
        return None

@st.cache_data(show_spinner="Loading data...")
def load_and_clean_data(file_path, mtime):
    raw_df = load_raw_sheet(file_path, mtime)
    df = raw_df.iloc[4:].copy()

    cluster_raw = df.iloc[:, 0]
    product_raw = df.iloc[:, 2]
    cluster = cluster_raw.astype(str).str.strip()
    product = product_raw.astype(str).str.strip()
    product_lower = product.str.lower()

    # صفوف المنتجات الحقيقية بس: لازم اسم منتج، ومش صف إجمالي ("... Total") ولا عنوان
    # (ملاحظة: مع pandas الجديد astype(str) بيسيب القيم الفاضية NaN، فلازم نفحص notna صراحة)
    is_product = (
        cluster_raw.notna() & product_raw.notna()
        & ~product_lower.isin(['', 'nan', 'none'])
        & ~cluster.str.lower().str.endswith('total')
        & ~product_lower.str.contains(r"^crm product name|\btotal\b|subtotal|\bsum\b", regex=True)
    )

    data = pd.DataFrame({
        'Cluster': cluster[is_product],
        'Product_ID': df.iloc[:, 1][is_product],
        'Product_Name': product[is_product],
    })

    for col_name, col_idx in NUMERIC_COLS.items():
        data[col_name] = pd.to_numeric(df.iloc[:, col_idx], errors='coerce').fillna(0)

    # الشيت بيحط 36525 (=100 سنة) في Stock Days لو مفيش مبيعات، فبنعتبرها "مفيش قيمة"
    data.loc[data['Stock_Days'] >= 36500, 'Stock_Days'] = float('nan')

    return data

@st.cache_data(show_spinner=False)
def build_line_workbook(file_path, mtime, line):
    """يبني ملف Excel فيه خط واحد بس (بيتكاشي لكل خط)."""
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    total_label = line + ' Total'

    rows_to_delete = []
    for row in range(5, ws.max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val is not None and str(val).strip() not in (line, total_label):
            rows_to_delete.append(row)

    for r in reversed(rows_to_delete):
        ws.delete_rows(r)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

LOGIN_CSS = """
<style>
/* ===== Login page only (this block is not rendered after sign-in) ===== */
[data-testid="stApp"] {
    background-image:
        radial-gradient(60% 50% at 15% 10%, rgba(230, 126, 34, 0.20), transparent 70%),
        radial-gradient(50% 45% at 90% 20%, rgba(52, 152, 219, 0.18), transparent 70%),
        radial-gradient(55% 50% at 60% 100%, rgba(155, 89, 182, 0.16), transparent 70%),
        linear-gradient(rgba(128,128,128,0.07) 1px, transparent 1px),
        linear-gradient(90deg, rgba(128,128,128,0.07) 1px, transparent 1px);
    background-size: auto, auto, auto, 44px 44px, 44px 44px;
}
[data-testid="stApp"]::before,
[data-testid="stApp"]::after {
    content: "";
    position: fixed;
    border-radius: 50%;
    filter: blur(70px);
    pointer-events: none;
    z-index: 0;
}
[data-testid="stApp"]::before {
    width: 420px; height: 420px; top: -120px; left: -100px;
    background: rgba(230, 126, 34, 0.35);
    animation: lg-float-a 14s ease-in-out infinite alternate;
}
[data-testid="stApp"]::after {
    width: 380px; height: 380px; bottom: -120px; right: -80px;
    background: rgba(52, 152, 219, 0.30);
    animation: lg-float-b 17s ease-in-out infinite alternate;
}
@keyframes lg-float-a { to { transform: translate(120px, 90px) scale(1.15); } }
@keyframes lg-float-b { to { transform: translate(-110px, -80px) scale(1.2); } }
@media (prefers-reduced-motion: reduce) {
    [data-testid="stApp"]::before, [data-testid="stApp"]::after { animation: none; }
}

[data-testid="stHeader"] { background: transparent; }

/* ===== Glass card ===== */
div[data-testid="stForm"] {
    position: relative;
    z-index: 1;
    max-width: 440px;
    margin: 0 auto;
    padding: 30px 30px 26px;
    border-radius: 24px;
    border: 1px solid color-mix(in srgb, currentColor 14%, transparent);
    background: color-mix(in srgb, currentColor 5%, transparent);
    -webkit-backdrop-filter: blur(18px) saturate(140%);
    backdrop-filter: blur(18px) saturate(140%);
    box-shadow:
        0 25px 70px -25px rgba(211, 84, 0, 0.45),
        0 0 60px -20px rgba(52, 152, 219, 0.35);
}
div[data-testid="stForm"]::before {
    content: "";
    position: absolute;
    top: 0; left: 12%; right: 12%;
    height: 2px;
    border-radius: 2px;
    background: linear-gradient(90deg, transparent, #e67e22, #3498db, transparent);
    box-shadow: 0 0 14px rgba(230, 126, 34, 0.8);
}

.lg-head { text-align: center; padding: 4px 0 18px; }
.lg-badge {
    width: 68px; height: 68px;
    margin: 0 auto 14px;
    display: flex; align-items: center; justify-content: center;
    border-radius: 20px;
    font-size: 32px;
    background: linear-gradient(135deg, rgba(230,126,34,0.25), rgba(52,152,219,0.20));
    border: 1px solid color-mix(in srgb, currentColor 16%, transparent);
    box-shadow: 0 0 28px rgba(230, 126, 34, 0.40);
    overflow: hidden;
}
.lg-badge img { width: 100%; height: 100%; object-fit: contain; padding: 8px; box-sizing: border-box; }
.lg-title {
    margin: 0 0 6px 0; padding: 0;
    font-size: 28px; font-weight: 700; letter-spacing: -0.3px; line-height: 1.15;
    background: linear-gradient(90deg, #d35400, #f39c12 55%, #f1c40f);
    -webkit-background-clip: text; background-clip: text;
    -webkit-text-fill-color: transparent; color: transparent;
}
.lg-sub { margin: 0 0 12px 0; font-size: 14px; opacity: 0.75; }
.lg-chip {
    display: inline-block;
    padding: 4px 12px;
    font-size: 11.5px; letter-spacing: 0.4px;
    border-radius: 999px;
    border: 1px solid color-mix(in srgb, currentColor 16%, transparent);
    background: color-mix(in srgb, currentColor 6%, transparent);
    opacity: 0.85;
}

/* ===== Inputs ===== */
div[data-testid="stForm"] div[data-baseweb="input"] {
    border-radius: 12px;
    border: 1px solid color-mix(in srgb, currentColor 24%, transparent);
    transition: border-color .2s ease, box-shadow .2s ease;
}
div[data-testid="stForm"] div[data-baseweb="input"]:focus-within {
    border-color: #e67e22;
    box-shadow: 0 0 0 3px rgba(230, 126, 34, 0.22), 0 0 22px rgba(230, 126, 34, 0.25);
}

/* ===== Sign-in button ===== */
div[data-testid="stFormSubmitButton"] button {
    height: 46px;
    border-radius: 12px;
    border: none;
    font-weight: 600;
    letter-spacing: 0.3px;
    color: #fff;
    background: linear-gradient(135deg, #f39c12, #d35400);
    box-shadow: 0 8px 24px -8px rgba(211, 84, 0, 0.8);
    transition: transform .15s ease, box-shadow .15s ease, filter .15s ease;
}
div[data-testid="stFormSubmitButton"] button:hover {
    color: #fff;
    transform: translateY(-1px);
    filter: brightness(1.08);
    box-shadow: 0 12px 30px -8px rgba(211, 84, 0, 0.95), 0 0 22px rgba(243, 156, 18, 0.45);
}
div[data-testid="stFormSubmitButton"] button:active { transform: translateY(0); }

@media (max-width: 640px) {
    div[data-testid="stForm"] { padding: 24px 18px 20px; border-radius: 20px; }
    .lg-title { font-size: 24px; }
}
</style>
"""

@st.cache_resource
def logo_data_uri():
    try:
        with open("logo.png", "rb") as f:
            return "data:image/png;base64," + base64.b64encode(f.read()).decode("ascii")
    except Exception:
        return None

def login_page():
    st.markdown(LOGIN_CSS, unsafe_allow_html=True)
    _, col2, _ = st.columns([1, 1.4, 1])
    with col2:
        st.markdown("<div style='height: 6vh;'></div>", unsafe_allow_html=True)

        with st.form("login_form"):
            logo = logo_data_uri()
            badge = f'<img src="{logo}" alt="logo">' if logo else "📊"
            st.markdown(f"""
                <div class="lg-head">
                    <div class="lg-badge">{badge}</div>
                    <h1 class="lg-title">&nbsp;Welcome</h1>
                    <p class="lg-sub">Sales Dashboard</p>
                    <span class="lg-chip">Commercial Excellence</span>
                </div>
            """, unsafe_allow_html=True)

            username = st.text_input("👤 Username", placeholder="Enter your username")
            password = st.text_input("🔑 Password", type="password", placeholder="Enter your password")
            st.markdown("<div style='margin-top: 15px;'></div>", unsafe_allow_html=True)
            submit = st.form_submit_button("Sign In", use_container_width=True)
            
            if submit:
                username = username.strip().lower()
                user = check_credentials(username, password)
                if user:
                    st.session_state.logged_in = True
                    st.session_state.username = username
                    st.session_state.role = user["role"]
                    st.session_state.assigned_line = user["line"]
                    st.rerun()
                else:
                    st.error("Incorrect username or password.")

if not st.session_state.logged_in:
    login_page()
else:
    # Sidebar Logo & User Info
    if os.path.exists("logo.png"):
        st.sidebar.image("logo.png", width=140)
    st.sidebar.title(f"👤 Welcome, {st.session_state.username}")
    if st.sidebar.button("Logout"):
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.session_state.role = ""
        st.session_state.assigned_line = ""
        st.rerun()

    st.title("📊 Sales & Distribution Dashboard")

    restore_shared_file()
    target_file = resolve_data_file()
    sheet_date = get_sheet_date(target_file) if target_file else None
    if sheet_date is not None:
        st.markdown(f"📅 **Date:** `{sheet_date}`")
    st.markdown("---")

    st.sidebar.header("📁 Control Panel")
    if sheet_date is not None:
        st.sidebar.info(f"📅 Date: {sheet_date}")

    flash = st.session_state.pop("flash", None)
    if flash:
        level, text = flash
        getattr(st.sidebar, level)(text)

    if st.session_state.role == "admin":
        # تغيير الـ key بيفضّي الـ uploader بعد كل رفع/مسح (وبيمنع تكرار المعالجة)
        uploader_n = st.session_state.get("uploader_n", 0)
        uploaded_file = st.sidebar.file_uploader(
            "Upload Excel File (Admin Only)", type=["xlsx", "xls"], key=f"uploader_{uploader_n}"
        )
        if uploaded_file is not None:
            content = uploaded_file.getvalue()
            with open(SHARED_FILE, "wb") as f:
                f.write(content)

            if github_cfg() is None:
                st.session_state.flash = ("warning", "Sheet uploaded, but storage is temporary (GitHub is not configured) and will be lost on restart.")
            else:
                with st.spinner("Saving to GitHub..."):
                    ok, msg = push_to_github(content, st.session_state.username)
                if ok:
                    st.session_state.flash = ("success", "Sheet uploaded and saved permanently.")
                else:
                    st.session_state.flash = ("warning", f"Sheet uploaded temporarily, but saving to GitHub failed: {msg}")

            restore_shared_file.clear()
            st.cache_data.clear()
            st.session_state.uploader_n = uploader_n + 1
            st.rerun()

        if target_file is not None:
            if st.sidebar.button("🗑️ Delete current sheet", use_container_width=True):
                st.session_state.confirm_delete = True

            if st.session_state.get("confirm_delete"):
                st.sidebar.warning("This will permanently delete the sheet for ALL users. Are you sure?")
                col_yes, col_no = st.sidebar.columns(2)
                if col_yes.button("Yes, delete", type="primary", use_container_width=True):
                    with st.spinner("Deleting..."):
                        ok, msg = delete_from_github(st.session_state.username)
                    if ok:
                        if os.path.exists(SHARED_FILE):
                            os.remove(SHARED_FILE)
                        st.cache_data.clear()
                        st.session_state.uploader_n = uploader_n + 1
                        st.session_state.flash = ("success", "Sheet deleted.")
                    else:
                        st.session_state.flash = ("error", f"Could not delete from GitHub, nothing was removed: {msg}")
                    st.session_state.confirm_delete = False
                    st.rerun()
                if col_no.button("Cancel", use_container_width=True):
                    st.session_state.confirm_delete = False
                    st.rerun()
    else:
        st.sidebar.info(f"👤 Line User: {st.session_state.assigned_line}")

    if target_file is not None:
        file_mtime = os.path.getmtime(target_file)
        df = load_and_clean_data(target_file, file_mtime)

        if st.session_state.role == "admin":
            st.sidebar.subheader("Filter Options (Admin)")
            selected_clusters = st.sidebar.multiselect(
                "Select Cluster / Line:",
                options=df['Cluster'].unique(),
                default=df['Cluster'].unique()
            )
            filtered_df = df[df['Cluster'].isin(selected_clusters)]

            st.sidebar.markdown("---")
            with open(target_file, "rb") as f:
                st.sidebar.download_button(
                    label="📥 Download All (Admin)",
                    data=f.read(),
                    file_name="Daily_Sales_All.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
        else:
            assigned_line = st.session_state.assigned_line
            filtered_df = df[df['Cluster'] == assigned_line]
            st.sidebar.warning(f"Viewing Line: {assigned_line}")

            if not filtered_df.empty:
                st.sidebar.markdown("---")
                st.sidebar.download_button(
                    label=f"📥 Download {assigned_line}",
                    data=build_line_workbook(target_file, file_mtime, assigned_line),
                    file_name=f"{assigned_line.replace(' ', '_')}_sales.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

        if filtered_df.empty:
            st.warning("⚠️ لا توجد بيانات مطابقة للخط الخاص بك في هذا الشيت. تأكد من تطابق اسم الخط (Cluster) مع اسم المستخدم.")
        else:
            st.subheader("📌 Performance Summary")
            
            total_sales = filtered_df['Total_Sales_Val'].sum()
            total_target = filtered_df['Total_Target_Val'].sum()
            total_ach = (total_sales / total_target * 100) if total_target > 0 else 0
            private_sales = filtered_df['Private_Sales_Val'].sum()
            tender_sales = filtered_df['Tender_Sales_Val'].sum()
            
            zero_target_count = (filtered_df['Total_Target_Val'] == 0).sum()
            total_products_count = len(filtered_df)
            
            k1, k2, k3, k4, k5, k6 = st.columns(6)
            k1.metric("Total Sales", format_currency(total_sales))
            k2.metric("Target", format_currency(total_target))
            k3.metric("Achievement %", f"{total_ach:.1f}%")
            k4.metric("Private Sales", format_currency(private_sales))
            k5.metric("Tender Sales", format_currency(tender_sales))
            k6.metric("Zero Target SKUs", f"{zero_target_count} / {total_products_count}")
            
            st.progress(min(max(int(total_ach), 0), 100), text=f"Overall Target Achievement: {total_ach:.1f}%")
            st.markdown("---")

            st.subheader("⚠️ Products Without Target (Zero Target Analysis)")
            no_target_df = filtered_df[filtered_df['Total_Target_Val'] == 0]
            
            if not no_target_df.empty:
                no_target_with_sales = no_target_df[no_target_df['Total_Sales_Val'] != 0]
                no_target_no_sales = no_target_df[no_target_df['Total_Sales_Val'] == 0]
                
                nt_col1, nt_col2 = st.columns(2)
                
                with nt_col1:
                    st.markdown(f"##### 🟢 No Target BUT Has Sales ({len(no_target_with_sales)} SKUs)")
                    if not no_target_with_sales.empty:
                        fig_nt_sales = px.bar(
                            no_target_with_sales.nlargest(10, 'Total_Sales_Val'),
                            x='Total_Sales_Val',
                            y='Product_Name',
                            orientation='h',
                            title="Top 10 Zero-Target Products by Sales",
                            text_auto=',.0f',
                            color_discrete_sequence=['#2ecc71']
                        )
                        fig_nt_sales.update_layout(yaxis={'categoryorder':'total ascending'})
                        st.plotly_chart(fig_nt_sales, use_container_width=True)
                    else:
                        st.info("No sales recorded for any zero-target product.")

                with nt_col2:
                    st.markdown(f"##### 🔴 No Target AND No Sales ({len(no_target_no_sales)} SKUs)")
                    st.dataframe(
                        style_numbers(no_target_no_sales[['Cluster', 'Product_Name', 'Total_Stock_Qty', 'Stock_Days']]),
                        use_container_width=True,
                        height=300
                    )

                with st.expander("📄 View All Products Without Target"):
                    st.dataframe(
                        style_numbers(no_target_df[['Cluster', 'Product_ID', 'Product_Name', 'Total_Sales_Val', 'Total_Stock_Qty', 'Stock_Days']].sort_values(by='Total_Sales_Val', ascending=False)),
                        use_container_width=True
                    )
            else:
                st.success("🎉 All selected products have an assigned target!")

            st.markdown("---")

            st.subheader("🌟 Top Product for Each Distributor")
            
            dist_names = DIST_NAMES
            top_products_data = []
            
            for dist in dist_names:
                dist_total = filtered_df[dist].sum()
                if dist_total > 0:
                    top_row = filtered_df.nlargest(1, dist).iloc[0]
                    top_products_data.append({
                        'Distributor': dist,
                        'Top Product': top_row['Product_Name'],
                        'Sales Value (EGP)': top_row[dist],
                        'Distributor Total Sales': dist_total,
                        '% of Distributor Sales': (top_row[dist] / dist_total * 100)
                    })
                    
            top_dist_df = pd.DataFrame(top_products_data)
            
            if not top_dist_df.empty:
                col_dist_cards, col_dist_selector = st.columns([3, 2])
                
                with col_dist_cards:
                    fig_top_dist = px.bar(
                        top_dist_df,
                        x='Sales Value (EGP)',
                        y='Distributor',
                        text='Top Product',
                        orientation='h',
                        title="Top Product Sales per Distributor",
                        color='Sales Value (EGP)',
                        color_continuous_scale='Blues'
                    )
                    fig_top_dist.update_traces(textposition='inside', texttemplate='%{text}: %{x:,.0f} EGP')
                    fig_top_dist.update_layout(yaxis={'categoryorder':'total ascending'}, showlegend=False)
                    st.plotly_chart(fig_top_dist, use_container_width=True)
                    
                with col_dist_selector:
                    st.markdown("##### 🔍 Deep Dive by Distributor")
                    selected_dist = st.selectbox("Choose a Distributor:", dist_names)
                    
                    top_5_dist = filtered_df.nlargest(5, selected_dist)[['Product_Name', selected_dist]]
                    top_5_dist = top_5_dist[top_5_dist[selected_dist] > 0]
                    
                    if not top_5_dist.empty:
                        fig_dist_deep = px.bar(
                            top_5_dist,
                            x=selected_dist,
                            y='Product_Name',
                            orientation='h',
                            title=f"Top 5 Products in {selected_dist}",
                            text_auto=',.0f',
                            color_discrete_sequence=['#2ecc71']
                        )
                        fig_dist_deep.update_layout(yaxis={'categoryorder':'total ascending'})
                        st.plotly_chart(fig_dist_deep, use_container_width=True)
                    else:
                        st.info(f"No sales recorded for {selected_dist} in selected filter(s).")
            else:
                st.info("No distributor sales available for current filter selection.")

            st.markdown("---")

            st.subheader("📈 Cluster Performance & Market Share")
            c1, c2 = st.columns(2)
            
            with c1:
                cluster_acc = filtered_df.groupby('Cluster').agg({
                    'Total_Sales_Val': 'sum',
                    'Total_Target_Val': 'sum'
                }).reset_index()
                # لو التارجت صفر منقسمش، ونكتب "No target" بدل 0% (عشان ميبانش إن الخط مقصّر)
                safe_target = cluster_acc['Total_Target_Val'].where(cluster_acc['Total_Target_Val'] > 0)
                cluster_acc['Ach %'] = (cluster_acc['Total_Sales_Val'] / safe_target * 100).round(1)
                cluster_acc['Label'] = cluster_acc['Ach %'].map(
                    lambda v: 'No target' if pd.isna(v) else f'Ach: {v:.1f}%'
                )
                cluster_acc['Ach_color'] = cluster_acc['Ach %'].fillna(0)

                fig_cluster = px.bar(
                    cluster_acc,
                    x='Cluster',
                    y='Total_Sales_Val',
                    text='Label',
                    title="Sales by Cluster (Label shows Achievement %)",
                    color='Ach_color',
                    color_continuous_scale='Greens',
                    labels={'Total_Sales_Val': 'Sales Value (EGP)', 'Ach_color': 'Ach %'}
                )
                fig_cluster.update_traces(textposition='outside')
                st.plotly_chart(fig_cluster, use_container_width=True)

            with c2:
                dist_totals = pd.DataFrame([
                    {'Distributor': d, 'Sales': filtered_df[d].sum()} for d in dist_names
                ])
                dist_totals = dist_totals[dist_totals['Sales'] > 0]
                
                if not dist_totals.empty:
                    fig_dist_pie = px.pie(
                        dist_totals,
                        names='Distributor',
                        values='Sales',
                        title="Distributor Share Breakdown",
                        hole=0.4,
                        color_discrete_sequence=px.colors.qualitative.Set2
                    )
                    fig_dist_pie.update_traces(
                        textinfo='percent+label',
                        texttemplate='%{label}<br>%{percent:.1%}',
                        insidetextorientation='radial'
                    )
                    st.plotly_chart(fig_dist_pie, use_container_width=True)
                    left_out = {d: filtered_df[d].sum() for d in dist_names if filtered_df[d].sum() < 0}
                    if left_out:
                        st.caption("Not shown (net returns): " + ", ".join(f"{d} ({v:,.0f} EGP)" for d, v in left_out.items()))

            st.markdown("---")

            st.subheader("🏆 Top Products & Biggest Target Gaps")
            p1, p2 = st.columns(2)
            
            with p1:
                top10_overall = filtered_df.nlargest(10, 'Total_Sales_Val')
                if not top10_overall.empty:
                    fig_top10 = px.bar(
                        top10_overall,
                        x='Total_Sales_Val',
                        y='Product_Name',
                        orientation='h',
                        title="Top 10 Products (Sales Value)",
                        text_auto=',.0f',
                        color_discrete_sequence=['#3498db']
                    )
                    fig_top10.update_layout(yaxis={'categoryorder':'total ascending'})
                    st.plotly_chart(fig_top10, use_container_width=True)

            with p2:
                gap_df = filtered_df[filtered_df['Total_Target_Val'] > 0].copy()
                gap_df['Target_Gap'] = gap_df['Total_Target_Val'] - gap_df['Total_Sales_Val']
                top_gap = gap_df[gap_df['Target_Gap'] > 0].nlargest(10, 'Target_Gap')
                if not top_gap.empty:
                    fig_gap = px.bar(
                        top_gap,
                        x='Target_Gap',
                        y='Product_Name',
                        orientation='h',
                        title="Top 10 Target Gap (Target − Sales)",
                        text_auto=',.0f',
                        color_discrete_sequence=['#e74c3c'],
                        labels={'Target_Gap': 'Gap to Target (EGP)'}
                    )
                    fig_gap.update_layout(yaxis={'categoryorder': 'total ascending'})
                    st.plotly_chart(fig_gap, use_container_width=True)

            st.markdown("---")

            st.subheader("📋 Product Sales Summary Table")
            search_q = st.text_input("🔍 Search product name:")
            if search_q:
                table_dir = filtered_df[filtered_df['Product_Name'].str.contains(search_q, case=False, na=False, regex=False)]
            else:
                table_dir = filtered_df

            disp_df = table_dir[[
                'Cluster', 'Product_Name', 'Total_Sales_Val', 
                'Total_Target_Val', 'Private_Sales_Val', 'Tender_Sales_Val', 'Stock_Days'
            ]].copy().sort_values(by='Total_Sales_Val', ascending=False)
            
            st.dataframe(
                style_numbers(disp_df),
                use_container_width=True,
                hide_index=True,
            )
    else:
        st.info("👋 يُرجى من الأدمن رفع ملف الشيت أولاً من لوحة التحكم ليتمكن باقي المستخدمين من استعراض البيانات.")
